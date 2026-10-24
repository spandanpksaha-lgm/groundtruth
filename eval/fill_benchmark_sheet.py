"""Write the team benchmark results back into the team's spreadsheet.

    python eval/fill_benchmark_sheet.py

Looks for the four 2x2 runs (chunking x embedding model) of the hard set:
    uc-hard-section220-main      uc-hard-section220-bgesmall
    uc-hard-fixed220-main        uc-hard-fixed220-bgesmall
Runs that don't exist yet are skipped, so it works with just the two MiniLM runs.

Reads eval/UC_GroundTruth_40_Question_Benchmark.xlsx (left untouched) and writes
eval/results/UC_benchmark_results.xlsx with:
  * the team's "Fixed-size result" and "Section-based result" columns filled in
    (MiniLM runs: PASS/FAIL + outcome label), one more result column per extra
    run, and each run's answer
  * a "Paired comparison" tab: the 2x2 grid of correct answers, totals per
    question type and primary stress, and the questions where systems disagree
"""
import argparse
import csv
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from groundtruth.grid import cell_label, cells, run_tag  # noqa: E402

RES = ROOT / "eval" / "results"
SHEET = ROOT / "eval" / "UC_GroundTruth_40_Question_Benchmark.xlsx"
GOOD = {"correct", "correct_refusal"}
PASS_FILL = PatternFill("solid", fgColor="E5F4E4")
FAIL_FILL = PatternFill("solid", fgColor="FBEBEB")
HEAD_FILL = PatternFill("solid", fgColor="EEF1F6")
F = "Arial"


def load(tag):
    path = RES / tag / "per_question.csv"
    return {r["id"]: r for r in csv.DictReader(open(path))} if path.exists() else None


def main(argv=None):
    p = argparse.ArgumentParser()
    p.add_argument("--fixed", help="override the fixed + MiniLM run tag")
    p.add_argument("--section", help="override the section + MiniLM run tag")
    p.add_argument("--out", default=str(RES / "UC_benchmark_results.xlsx"))
    a = p.parse_args(argv)

    # (short name, chunker, model label, tag, results)
    runs = []
    for ch, model, _ in cells("uc"):
        tag = run_tag("uc", "hard", ch, model)
        if ch == "fixed" and a.fixed and tag.endswith("-main"):
            tag = a.fixed
        if ch == "section" and a.section and tag.endswith("-main"):
            tag = a.section
        res = load(tag)
        if res is None:
            print(f"  (skipping {tag}: not run yet)")
            continue
        runs.append((f"{ch.capitalize()} + {cell_label(model)}", ch, cell_label(model), tag, res))
    if not runs:
        sys.exit("No hard-set runs found. Run eval/run_eval.py --corpus uc --split hard first.")

    wb = load_workbook(SHEET)
    ws = wb["UC 40 Questions"]
    head = [c.value for c in ws[1]]
    team_cols = {("fixed", "MiniLM"): head.index("Fixed-size result") + 1,
                 ("section", "MiniLM"): head.index("Section-based result") + 1}
    nxt = ws.max_column + 1
    result_col, answer_col = {}, {}
    for name, ch, ml, tag, _ in runs:
        if (ch, ml) in team_cols:
            result_col[name] = team_cols[(ch, ml)]
        else:
            result_col[name] = nxt
            c = ws.cell(1, nxt, f"{name} result")
            c.font = Font(name=F, bold=True)
            ws.column_dimensions[c.column_letter].width = 26
            nxt += 1
    for name, *_ in runs:
        answer_col[name] = nxt
        c = ws.cell(1, nxt, f"{name} answer")
        c.font = Font(name=F, bold=True)
        ws.column_dimensions[c.column_letter].width = 55
        nxt += 1

    stats = defaultdict(lambda: defaultdict(int))   # group -> run name -> correct
    count = defaultdict(int)
    disagree = []
    for r in range(2, ws.max_row + 1):
        qid = ws.cell(r, 1).value
        if not qid or any(qid not in res for *_, res in runs):
            continue
        oks = {}
        for name, _, _, _, res in runs:
            row = res[qid]
            ok = row["outcome"] in GOOD
            oks[name] = ok
            c = ws.cell(r, result_col[name], f"{'PASS' if ok else 'FAIL'} ({row['outcome']})")
            c.fill = PASS_FILL if ok else FAIL_FILL
            c.font = Font(name=F, size=10)
            c = ws.cell(r, answer_col[name], row["answer"])
            c.font = Font(name=F, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for key in ("ALL", f"Type: {ws.cell(r, 2).value}", f"Primary stress: {ws.cell(r, 7).value}"):
            count[key] += 1
            for name, ok in oks.items():
                stats[key][name] += ok
        if len(set(oks.values())) > 1:
            disagree.append([qid, ws.cell(r, 4).value]
                            + [f"{'PASS' if oks[n] else 'FAIL'} ({res[qid]['outcome']})"
                               for n, _, _, _, res in runs])

    if "Paired comparison" in wb.sheetnames:
        del wb["Paired comparison"]
    pc = wb.create_sheet("Paired comparison")
    bold_rows = []
    n_all = count["ALL"]

    # 1) the 2x2 grid
    models = list(dict.fromkeys(ml for _, _, ml, _, _ in runs))
    by = {(ch, ml): stats["ALL"][name] for name, ch, ml, _, _ in runs}
    pc.append([f"Correct out of {n_all} (PASS = right answer with a supporting citation, or a correct refusal)"])
    bold_rows.append(pc.max_row)
    pc.append(["Chunking \\ Embedding"] + models)
    bold_rows.append(pc.max_row)
    for ch in ("section", "fixed"):
        pc.append([ch.capitalize()] + [f"{by[(ch, m)]}/{n_all}" if (ch, m) in by else "not run" for m in models])
    pc.append([])

    # 2) breakdown by group
    names = [n for n, *_ in runs]
    pc.append(["Group", "Questions"] + names)
    bold_rows.append(pc.max_row)
    order = ["ALL"] + sorted(k for k in count if k.startswith("Primary")) + sorted(k for k in count if k.startswith("Type"))
    for k in order:
        pc.append([k, count[k]] + [stats[k][n] for n in names])
    pc.append([])

    # 3) disagreements
    pc.append([f"Questions where the systems disagree ({len(disagree)})"])
    bold_rows.append(pc.max_row)
    pc.append(["ID", "Question"] + names)
    bold_rows.append(pc.max_row)
    for row in disagree:
        pc.append(row)
    pc.append([])
    pc.append(["Runs: " + "; ".join(f"{n} = {t}" for n, _, _, t, _ in runs)
               + ". Small sample: a 1-2 question difference is not significant."])
    for row in pc.iter_rows():
        for c in row:
            c.font = Font(name=F, size=10, bold=c.row in bold_rows)
            if c.row in bold_rows[1:2] or c.row in bold_rows[2:3] or c.row in bold_rows[4:5]:
                c.fill = HEAD_FILL
            if isinstance(c.value, str) and c.value.startswith(("PASS", "FAIL")):
                c.fill = PASS_FILL if c.value.startswith("PASS") else FAIL_FILL
    for col, w in zip("ABCDEF", (34, 60, 24, 24, 24, 24)):
        pc.column_dimensions[col].width = w
    wb.save(a.out)
    print("Correct out of", n_all, "->", ", ".join(f"{n}: {stats['ALL'][n]}" for n in names))
    print(f"{len(disagree)} questions where systems disagree. Wrote {a.out}")


if __name__ == "__main__":
    main()
