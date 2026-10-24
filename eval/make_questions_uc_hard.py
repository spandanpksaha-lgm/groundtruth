"""Team-written hard benchmark for the UC Student Handbook (40 questions).

Source of truth: eval/UC_GroundTruth_40_Question_Benchmark.xlsx (written by the
team, before seeing model output). This script keeps the team's question,
type, expected answer and page, and adds what automated scoring needs:

  evidence    exact quotes from the handbook (the sheet gives page numbers; the
              sheet's own guide asks for exact quotations on import)
  key_facts   words a correct answer must contain (groups = all required,
              alternatives within a group = any one); the LLM judge
              (`--judge`) gives a second opinion on wording-sensitive rows

    python eval/make_questions_uc_hard.py    # -> eval/questions_uc_hard.jsonl
    python eval/verify_questions.py --corpus uc --split hard
"""
import json
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).parent
SHEET = HERE / "UC_GroundTruth_40_Question_Benchmark.xlsx"

NO = ["no,", "no.", "not ", "isn't", "doesn't", "does not", "do not", "cannot", "can't", "incorrect"]

# id -> (key_facts, evidence, stakes). Added by Claude; questions and answers are the team's.
SCORING = {
    "UC01": ([["jeopardize", "not automatic", "not definitely", "doesn't say", "does not say", "does not state",
               "not stated", "may "]], ["Lack of participation during this time may jeopardize enrollment status"], "medium"),
    "UC02": ([["catalog"]], ["graduate university catalog for academic standing policy"], "high"),
    "UC03": ([["instructor"], ["chair"], ["vice president for academic affairs", "vpaa"]],
             ["must discuss the grade with", "formal written appeal with the Vice"], "high"),
    "UC04": ([["last day of class"]], ["must be filed four weeks after the last day of class"], "high"),
    "UC05": ([["“w”", "\"w\"", "'w'", "grade of w", "a w ", "w grade", "mark of w", "w for", " w,", " w."],
              ["professor", "instructor"]],
             ["carries no grade point penalty", "receive a grade as determined by the professor"], "high"),
    "UC06": ([NO + ["normally"]], ["partial medical/ emergency withdrawals are not permitted"], "high"),
    "UC07": ([["three weeks", "3 weeks"]], ["within three weeks of becoming unable to attend classes"], "high"),
    "UC08": ([["documentation", "letter", "medical professional"]],
             ["must be supported by a letter from a medical professional"], "medium"),
    "UC09": ([["25", "twenty-five"]], ["complete and report 200+ community service hours"], "low"),
    "UC10": ([["automatic"]], ["will receive 200 community service hours automatically"], "low"),
    "UC11": ([["attendance", "attending", "attend"]], ["Confirmation of Enrollment normally occurs by attendance"], "medium"),
    "UC12": ([["included", "listed", "is directory", "are directory", "not excluded", "counts as", "is an example",
               "example of directory", "examples of directory", "part of the directory"]],
             ["addresses including electronic mail address"], "medium"),
    "UC13": ([["october 19", "oct 19", "oct. 19", "10/19"], ["october 27", "oct 27", "oct. 27", "10/27"]],
             ["Monday, October 19, 2026", "Tuesday, October 27, 2026"], "medium"),
    "UC14": ([["april 23"], ["april 28"]], ["Friday, April 23, 2027", "April 28, 2027"], "medium"),
    "UC15": ([NO], ["Campus-issued temporary permits do not authorize parking in state-marked"], "medium"),
    "UC16": ([["head athletic trainer", "athletic trainer"], ["public safety", "parking control", "travel services"]],
             ["must obtain their permit through the Head Athletic Trainer"], "medium"),
    "UC17": ([["annual", "every year", "each year", "yearly", "one academic year"],
              ["non-transferrable", "non-transferable", "nontransferable", "not transferable", "not transferrable",
               "cannot be transferred", "can't be transferred"]],
             ["Permits are non-transferrable", "registered online annually"], "low"),
    "UC18": ([["august 15"], ["august 14"]], ["August 15th through August 14th"], "low"),
    "UC19": ([["5:00", "5 pm", "5pm", "5 p.m", "five"]],
             ["General lots are for visitors and commuter students only"], "low"),
    "UC20": ([["moss"]], ["always restricted to Moss residents"], "low"),
    "UC21": ([["tow"]], ["(4) or more citations within (1) semester will be towed"], "medium"),
    "UC22": ([["towing"], ["65"]], ["Cost to retrieve towed vehicles begins at $65.00"], "low"),
    "UC23": ([["three nights", "3 nights"]], ["minimum of three nights per"], "high"),
    "UC24": ([["100"]], ["charged a fee at the rate of $20 per day"], "low"),
    "UC25": ([["150"]], ["assessed a minimum $50.00 fine", "minimum of $50.00 per item"], "medium"),
    "UC26": ([["30", "thirty"]], ["minimum of 30 credit hours completed"], "medium"),
    "UC27": ([["housing application"]], ["Housing Application for the same term"], "medium"),
    "UC28": ([["24 hours", "24-hour", "twenty-four"]], ["within 24 hours of its receipt"], "high"),
    "UC29": ([["300"]], ["minimum of $100 per day while the animal"], "medium"),
    "UC30": ([["apply", "followed", "still", "also", "must follow", "required to follow"]],
             ["all Residence Life policies and procedures stated in the Student"], "medium"),
    "UC31": ([NO + ["prohibited", "not permitted", "not allowed"], ["coffee maker"]],
             ["crock pots, air fryers, and rice", "Coffee makers, blow dryers, hair curlers"], "medium"),
    "UC32": ([NO + ["prohibited", "not permitted", "not allowed"], ["residence life"]],
             ["Students are not permitted to use their own pest control products"], "low"),
}
UNANSWERABLE_STAKES = {"UC33": "low", "UC34": "low", "UC35": "medium", "UC36": "low", "UC37": "high",
                       "UC38": "low", "UC39": "low", "UC40": "high"}


def main():
    ws = load_workbook(SHEET, data_only=True)["UC 40 Questions"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    out = []
    for qid, qtype, should, question, expected, location, stress, *_ in rows:
        if not qid:
            continue
        answerable = str(should).strip().lower().startswith("y")
        base = dict(id=qid, question=question.strip(), answerable=answerable,
                    category=str(qtype).strip().lower().replace(" / ", "_").replace(" ", "_"),
                    reference=expected.strip(), source=location, primary_stress=stress,
                    author="team")
        if answerable:
            key_facts, evidence, stakes = SCORING[qid]
            base.update(key_facts=key_facts, evidence=evidence, stakes=stakes)
        else:
            base.update(category="unanswerable_team", stakes=UNANSWERABLE_STAKES.get(qid, "medium"))
        out.append(base)
    path = HERE / "questions_uc_hard.jsonl"
    with open(path, "w") as f:
        for r in out:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    print(f"wrote {len(out)} -> {path.name} ({sum(not r['answerable'] for r in out)} unanswerable)")


if __name__ == "__main__":
    main()
